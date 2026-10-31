"""Import subject marks from CSV/XLSX and export calculated results.

Expected input columns:
student_name, semester, subject, marks, credits
Optional column: register_number

One row represents one subject for one student in one semester.
"""
from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from pathlib import Path
from typing import Any

try:
    from .gradecalc import calculate_sgpa
except ImportError:
    from gradecalc import calculate_sgpa


REQUIRED_COLUMNS = {"student_name", "semester", "subject", "marks", "credits"}
OUTPUT_COLUMNS = [
    "student_name", "register_number", "semester", "subject", "marks",
    "grade", "grade_point", "credits", "credit_points", "semester_sgpa",
]


def _clean_row(row: dict[str, Any], row_number: int) -> dict[str, Any]:
    normalized = {str(key).strip().lower(): value for key, value in row.items() if key is not None}
    missing = REQUIRED_COLUMNS - normalized.keys()
    if missing:
        raise ValueError(f"Row {row_number}: missing columns: {', '.join(sorted(missing))}")
    student = str(normalized.get("student_name") or "").strip()
    semester = str(normalized.get("semester") or "").strip()
    subject = str(normalized.get("subject") or "").strip()
    if not student or not semester or not subject:
        raise ValueError(f"Row {row_number}: student_name, semester, and subject cannot be blank")
    try:
        credits = float(normalized["credits"])
    except (TypeError, ValueError):
        raise ValueError(f"Row {row_number}: credits must be numeric") from None
    if credits <= 0:
        raise ValueError(f"Row {row_number}: credits must be greater than zero")
    marks = normalized.get("marks")
    if marks is not None:
        marks = str(marks).strip()
    return {
        "student_name": student,
        "register_number": str(normalized.get("register_number") or "").strip(),
        "semester": semester,
        "subject": subject,
        "marks": marks if marks else "",
        "credits": credits,
    }


def load_records(input_path: str | Path) -> list[dict[str, Any]]:
    """Load and validate records from .csv or .xlsx."""
    path = Path(input_path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            if reader.fieldnames is None:
                raise ValueError("CSV file must include a header row")
            return [_clean_row(row, index) for index, row in enumerate(reader, start=2)]
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Install dependencies with: python -m pip install -r python/requirements.txt") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            headers = next(iterator, None)
            if not headers:
                raise ValueError("Excel file must include a header row")
            names = [str(value).strip().lower() if value is not None else "" for value in headers]
            records = []
            for row_number, values in enumerate(iterator, start=2):
                if all(value is None for value in values):
                    continue
                records.append(_clean_row(dict(zip(names, values)), row_number))
            return records
        finally:
            workbook.close()
    raise ValueError("Input file must use .csv or .xlsx")


def calculate_results(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Calculate each student's SGPA for each semester and return flat result rows."""
    groups: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        key = (
            record["student_name"],
            record["register_number"],
            record["semester"],
        )
        groups[key].append(record)

    output = []
    for (student, register, semester), subjects in sorted(groups.items()):
        calculated = calculate_sgpa([
            {"name": item["subject"], "marks": item["marks"], "credits": item["credits"]}
            for item in subjects
        ])
        if calculated is None:
            raise ValueError(f"No valid subject credits for {student}, {semester}")
        grade_by_name = {item["name"]: item for item in calculated["validSubjects"]}
        for item in subjects:
            grade = grade_by_name[item["subject"]]
            output.append({
                "student_name": student,
                "register_number": register,
                "semester": semester,
                "subject": item["subject"],
                "marks": item["marks"],
                "grade": grade["gradeLabel"],
                "grade_point": grade["gradePoint"],
                "credits": item["credits"],
                "credit_points": item["numericCredits"] * grade["gradePoint"],
                "semester_sgpa": calculated["sgpa"],
            })
    return output


def export_results(rows: list[dict[str, Any]], output_path: str | Path) -> Path:
    """Write calculated rows to CSV or XLSX."""
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=OUTPUT_COLUMNS)
            writer.writeheader()
            writer.writerows(rows)
        return path
    if suffix == ".xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise RuntimeError("Install dependencies with: python -m pip install -r python/requirements.txt") from exc
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Subject Results"
        sheet.append(OUTPUT_COLUMNS)
        for row in rows:
            sheet.append([row.get(column, "") for column in OUTPUT_COLUMNS])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_cells in sheet.columns:
            width = min(max(max(len(str(cell.value or "")) for cell in column_cells) + 2, 12), 32)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
        workbook.save(path)
        return path
    raise ValueError("Output file must use .csv or .xlsx")


def main() -> None:
    parser = argparse.ArgumentParser(description="Calculate and export GradeCalc student results.")
    parser.add_argument("input_file", help="Input .csv or .xlsx with subject records")
    parser.add_argument("output_file", help="Output .csv or .xlsx")
    args = parser.parse_args()
    records = load_records(args.input_file)
    results = calculate_results(records)
    destination = export_results(results, args.output_file)
    print(f"Processed {len(records)} subject records for {len({(r['student_name'], r['semester']) for r in records})} student-semester groups.")
    print(f"Results exported to: {destination}")


if __name__ == "__main__":
    main()
